#!/usr/bin/env python3
"""Build a Word document (.docx) from Markdown, for technical documentation that is shared as Word.

Supports what technical documentation uses: headings (#..####), paragraphs, bulleted and numbered
lists (nested by two spaces), tables (| a | b |), fenced code blocks, block quotes, and inline
**bold**, *italic* and `code`. With --template the styles (fonts, colours, heading numbering,
header / footer) of an existing .docx are used and its body is replaced.

Requires python-docx (`python -m pip install python-docx`).

Usage:
  python md_to_docx.py <doc.md> [<out.docx>] [--template <company.docx>]
  without <out.docx> the document goes next to the Markdown file (<name>.docx).
"""
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
except ImportError:
    sys.exit("python-docx is missing: python -m pip install python-docx")

INLINE = re.compile(r"(\*\*(.+?)\*\*|`([^`]+)`|(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*]))")
CODE_FONT = "Consolas"


def add_inline(paragraph, text):
    position = 0
    for match in INLINE.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        if match.group(2) is not None:
            paragraph.add_run(match.group(2)).bold = True
        elif match.group(3) is not None:
            run = paragraph.add_run(match.group(3))
            run.font.name = CODE_FONT
        else:
            paragraph.add_run(match.group(4)).italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def shade(element, colour):
    properties = element.get_or_add_tcPr() if element.tag.endswith("}tc") else element.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), colour)
    properties.append(shading)


def style_or(document, *names):
    """The first existing paragraph or table style, as an object: looking a style up by name fails
    in python-docx when the template spells a built-in name differently ("Heading 2")."""
    available = {s.name.lower(): s for s in document.styles if s.type in (1, 3)}
    return next((available[n.lower()] for n in names if n.lower() in available), None)


def cells(line):
    line = line.strip().strip("|")
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line)]


def add_table(document, rows):
    header, body = rows[0], [r for r in rows[1:] if not re.fullmatch(r"\s*:?-{2,}:?\s*", "".join(r))]
    width = max(len(header), *(len(r) for r in body)) if body else len(header)
    table = document.add_table(rows=1 + len(body), cols=width)
    grid = style_or(document, "Table Grid")
    if grid:
        table.style = grid
    for row_index, row in enumerate([header] + body):
        for col in range(width):
            cell = table.cell(row_index, col)
            paragraph = cell.paragraphs[0]
            for part_index, part in enumerate((row[col] if col < len(row) else "").split("<br>")):
                if part_index:
                    paragraph.add_run().add_break(WD_BREAK.LINE)
                add_inline(paragraph, part)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
                shade(cell._tc, "D9E2F3")
    document.add_paragraph()


def add_code(document, lines):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    shade(paragraph._p, "F2F2F2")
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        run = paragraph.add_run(line or " ")
        run.font.name = CODE_FONT
        run.font.size = Pt(8.5)


def add_heading(document, text, level):
    name = style_or(document, "Title" if level == 0 else f"Heading {level}")
    if name:
        add_inline(document.add_paragraph(style=name), text)
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    add_inline(paragraph, text)
    for run in paragraph.runs:
        run.bold = True
        run.font.size = Pt({0: 20, 1: 16, 2: 14, 3: 12}.get(level, 11))


def build(markdown, document):
    lines = markdown.splitlines()
    bullet = style_or(document, "List Bullet")
    number = style_or(document, "List Number")
    quote = style_or(document, "Quote", "Intense Quote")
    index = 0
    paragraph_lines = []

    def flush():
        if paragraph_lines:
            add_inline(document.add_paragraph(), " ".join(paragraph_lines))
            paragraph_lines.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            block = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            add_code(document, block)
        elif re.match(r"^#{1,6}\s", stripped):
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            add_heading(document, text, 0 if level == 1 and not document.paragraphs else min(level, 4))
        elif stripped.startswith("|"):
            flush()
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(cells(lines[index]))
                index += 1
            add_table(document, rows)
            continue
        elif re.match(r"^\s*([-*+]|\d+[.)])\s+", line):
            flush()
            indent = (len(line) - len(line.lstrip())) // 2
            ordered = re.match(r"^\s*\d+[.)]\s+", line) is not None
            text = re.sub(r"^\s*([-*+]|\d+[.)])\s+", "", line)
            name = number if ordered else bullet
            if name and indent:
                name = style_or(document, f"{name.name} {min(indent + 1, 3)}") or name
            paragraph = document.add_paragraph(style=name)
            if not name:
                paragraph.add_run("- ")
            add_inline(paragraph, text)
        elif stripped.startswith(">"):
            flush()
            paragraph = document.add_paragraph(style=quote)
            add_inline(paragraph, stripped.lstrip("> ").strip())
            if not quote:
                for run in paragraph.runs:
                    run.italic = True
                    run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        elif re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped) or not stripped:
            flush()
        else:
            paragraph_lines.append(stripped)
        index += 1
    flush()


def main(argv):
    args = list(argv)
    template = None
    if "--template" in args:
        position = args.index("--template")
        template = args[position + 1]
        del args[position:position + 2]
    if not args:
        print(__doc__)
        return 2
    source = Path(args[0])
    target = Path(args[1]) if len(args) > 1 else source.with_suffix(".docx")
    document = Document(template) if template else Document()
    if template:
        body = document.element.body
        for child in list(body):
            if not child.tag.endswith("}sectPr"):
                body.remove(child)
    build(source.read_text(encoding="utf-8"), document)
    document.save(target)
    print(f"written {target}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
